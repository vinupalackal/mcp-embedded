from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

out_file = "/Users/vpalac201@cable.comcast.com/Project/rdk-service-hub/docs/rdk-service-hub-customer-brief.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BLUE = RGBColor(0x1F, 0x6F, 0xB6)
TEXT = RGBColor(0x1E, 0x2A, 0x39)
ACCENT = RGBColor(0x27, 0x9E, 0xC2)
LIGHT = RGBColor(0xF7, 0xF9, 0xFC)
GREEN = RGBColor(0x20, 0x9B, 0x5A)


def add_title_slide(title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
    strip = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.22))
    strip.fill.solid(); strip.fill.fore_color.rgb = NAVY
    strip.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(8.1), Inches(0.9))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.75), Inches(2.1), Inches(11.0), Inches(0.8))
        sp = sub.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(16)
        sp.font.color.rgb = TEXT

    card = slide.shapes.add_shape(1, Inches(9.3), Inches(1.6), Inches(2.8), Inches(1.8))
    card.fill.solid(); card.fill.fore_color.rgb = BLUE
    card.line.fill.background()
    card.text = "Customer\nBrief"
    card.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    card.text_frame.paragraphs[0].font.size = Pt(20)
    card.text_frame.paragraphs[0].font.bold = True
    card.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def add_content_slide(title, bullets, accent_color=BLUE, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT

    strip = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.55))
    strip.fill.solid(); strip.fill.fore_color.rgb = accent_color
    strip.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.8), Inches(8.2), Inches(0.45))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY

    panel = slide.shapes.add_shape(1, Inches(0.7), Inches(1.5), Inches(11.8), Inches(4.8))
    panel.fill.solid(); panel.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    panel.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    panel.line.width = Pt(1)

    tf = panel.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.15)
    tf.margin_bottom = Inches(0.15)

    for i, bullet in enumerate(bullets):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.text = bullet
        par.level = 0
        par.bullet = True
        par.font.size = Pt(20)
        par.font.color.rgb = TEXT
        par.space_after = Pt(8)

    if note:
        note_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.45))
        np = note_box.text_frame.paragraphs[0]
        np.text = note
        np.font.size = Pt(11)
        np.font.color.rgb = RGBColor(0x4A, 0x5A, 0x68)


def add_two_col_slide(title, left, right, accent_color=BLUE):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT

    strip = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.55))
    strip.fill.solid(); strip.fill.fore_color.rgb = accent_color
    strip.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.8), Inches(8.0), Inches(0.45))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY

    left_box = slide.shapes.add_shape(1, Inches(0.7), Inches(1.5), Inches(5.4), Inches(4.9))
    left_box.fill.solid(); left_box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    left_box.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    left_box.line.width = Pt(1)
    ltf = left_box.text_frame; ltf.word_wrap = True
    ltf.margin_left = Inches(0.25); ltf.margin_right = Inches(0.2)
    for i, item in enumerate(left):
        par = ltf.paragraphs[0] if i == 0 else ltf.add_paragraph()
        par.text = item
        par.bullet = True
        par.font.size = Pt(18)
        par.font.color.rgb = TEXT
        par.space_after = Pt(8)

    right_box = slide.shapes.add_shape(1, Inches(6.6), Inches(1.5), Inches(5.9), Inches(4.9))
    right_box.fill.solid(); right_box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    right_box.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    right_box.line.width = Pt(1)
    rtf = right_box.text_frame; rtf.word_wrap = True
    rtf.margin_left = Inches(0.25); rtf.margin_right = Inches(0.2)
    for i, item in enumerate(right):
        par = rtf.paragraphs[0] if i == 0 else rtf.add_paragraph()
        par.text = item
        par.bullet = True
        par.font.size = Pt(18)
        par.font.color.rgb = TEXT
        par.space_after = Pt(8)


def add_summary_slide(title, lines):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
    strip = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.55))
    strip.fill.solid(); strip.fill.fore_color.rgb = ACCENT
    strip.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.8), Inches(9.0), Inches(0.45))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY

    panel = slide.shapes.add_shape(1, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.9))
    panel.fill.solid(); panel.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    panel.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    panel.line.width = Pt(1)
    tf = panel.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.2)
    for i, line in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.text = line
        par.font.size = Pt(18)
        par.font.color.rgb = TEXT
        par.space_after = Pt(10)


add_title_slide(
    "Multi-Plane Runtime Manager",
    "Customer-facing overview of the diagnostic and management execution platform"
)

add_content_slide(
    "The challenge",
    [
        "Device operations need fast, reliable diagnostics across multiple execution planes.",
        "Different actions require different levels of safety, control, and operational visibility.",
        "A single runtime must support both routine diagnostics and higher-risk management workflows.",
        "The system must remain secure while still being extensible and easy to manage."
    ],
    accent_color=BLUE,
    note="The solution needs to balance flexibility, policy control, and operational trust."
)

add_content_slide(
    "What the platform provides",
    [
        "A single runtime that routes requests by plane: triage, management, control, and config-apply.",
        "Support for static tool catalogs and optional dynamic plugin extensions.",
        "Command policy checks to reject unsafe or destructive operations before execution.",
        "Resource guardrails for CPU, memory, disk, and process pressure.",
        "Structured telemetry and logs for traceability and troubleshooting."
    ],
    accent_color=GREEN,
    note="This gives operators a consistent interface for diagnostics and device-side operations."
)

add_two_col_slide(
    "How it works",
    [
        "A request is received from a cloud or management system.",
        "The runtime identifies the correct plane and tool.",
        "The tool is looked up in a catalog or plugin registry.",
        "The request is validated before any command is launched."
    ],
    [
        "Resource checks verify host health and process pressure.",
        "Command execution is bounded by timeout and output handling.",
        "The system returns execution status and output to the caller.",
        "Logs and manifest data keep the runtime observable and auditable."
    ],
    accent_color=NAVY
)

add_content_slide(
    "Why customers value it",
    [
        "Consistent execution across multiple tool planes without separate code paths.",
        "Policy-aware execution reduces operational risk and protects against unsafe commands.",
        "Flexible deployment model with configurable files and environment overrides.",
        "Extensible design supports future tool additions with minimal runtime disruption.",
        "Operational clarity through logging, status reporting, and runtime guardrails."
    ],
    accent_color=ACCENT,
    note="The result is a reliable, governed platform for diagnostics and control workflows."
)

add_summary_slide(
    "Customer takeaway",
    [
        "Multi-plane Runtime Manager delivers a secure, scalable, and extensible execution model for diagnostics and operational support.",
        "It unifies tool routing, runtime protection, and operational observability under one managed platform.",
        "This enables safer device operations while preserving flexibility for future capabilities and deployments."
    ]
)

prs.save(out_file)
print(f"Created presentation: {out_file}")
