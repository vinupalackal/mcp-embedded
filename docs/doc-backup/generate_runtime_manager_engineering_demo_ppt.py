from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

out_file = "/Users/vpalac201@cable.comcast.com/Project/rdk-service-hub/docs/rdk-service-hub-engineering-demo.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BLUE = RGBColor(0x1F, 0x6F, 0xB6)
TEXT = RGBColor(0x1E, 0x2A, 0x39)
ACCENT = RGBColor(0x27, 0x9E, 0xC2)
LIGHT = RGBColor(0xF7, 0xF9, 0xFC)
GREEN = RGBColor(0x20, 0x9B, 0x5A)


def add_title_slide(title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
    strip = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.22))
    strip.fill.solid(); strip.fill.fore_color.rgb = NAVY
    strip.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(8.4), Inches(0.9))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY

    sub = slide.shapes.add_textbox(Inches(0.75), Inches(2.1), Inches(11.0), Inches(0.8))
    sp = sub.text_frame.paragraphs[0]
    sp.text = subtitle
    sp.font.size = Pt(16)
    sp.font.color.rgb = TEXT

    card = slide.shapes.add_shape(1, Inches(9.4), Inches(1.6), Inches(2.8), Inches(1.8))
    card.fill.solid(); card.fill.fore_color.rgb = BLUE
    card.line.fill.background()
    card.text = "Engineering\nDemo"
    card.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    card.text_frame.paragraphs[0].font.size = Pt(18)
    card.text_frame.paragraphs[0].font.bold = True
    card.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def add_content_slide(title, bullets, accent_color=BLUE, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT

    strip = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.55))
    strip.fill.solid(); strip.fill.fore_color.rgb = accent_color
    strip.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.8), Inches(8.0), Inches(0.45))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY

    panel = slide.shapes.add_shape(1, Inches(0.7), Inches(1.5), Inches(11.8), Inches(4.8))
    panel.fill.solid(); panel.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    panel.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    panel.line.width = Pt(1)

    tf = panel.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.15); tf.margin_bottom = Inches(0.15)
    for i, bullet in enumerate(bullets):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.text = bullet
        par.level = 0
        par.bullet = True
        par.font.size = Pt(19)
        par.font.color.rgb = TEXT
        par.space_after = Pt(8)

    if note:
        note_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.45))
        np = note_box.text_frame.paragraphs[0]
        np.text = note
        np.font.size = Pt(11)
        np.font.color.rgb = RGBColor(0x4A, 0x5A, 0x68)


def add_two_col_slide(title, left, right, accent_color=BLUE):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT

    strip = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.55))
    strip.fill.solid(); strip.fill.fore_color.rgb = accent_color
    strip.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.8), Inches(8.0), Inches(0.45))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY

    left_box = slide.shapes.add_shape(1, Inches(0.7), Inches(1.55), Inches(5.5), Inches(4.95))
    left_box.fill.solid(); left_box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    left_box.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    left_box.line.width = Pt(1)
    ltf = left_box.text_frame; ltf.word_wrap = True
    ltf.margin_left = Inches(0.25); ltf.margin_right = Inches(0.2)
    for i, item in enumerate(left):
        par = ltf.paragraphs[0] if i == 0 else ltf.add_paragraph()
        par.text = item
        par.bullet = True
        par.font.size = Pt(18)
        par.font.color.rgb = TEXT
        par.space_after = Pt(8)

    right_box = slide.shapes.add_shape(1, Inches(6.7), Inches(1.55), Inches(5.9), Inches(4.95))
    right_box.fill.solid(); right_box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    right_box.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    right_box.line.width = Pt(1)
    rtf = right_box.text_frame; rtf.word_wrap = True
    rtf.margin_left = Inches(0.25); rtf.margin_right = Inches(0.2)
    for i, item in enumerate(right):
        par = rtf.paragraphs[0] if i == 0 else rtf.add_paragraph()
        par.text = item
        par.bullet = True
        par.font.size = Pt(18)
        par.font.color.rgb = TEXT
        par.space_after = Pt(8)


def add_diag_slide(title, lines):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
    strip = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.55))
    strip.fill.solid(); strip.fill.fore_color.rgb = ACCENT
    strip.line.fill.background()
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.8), Inches(8.5), Inches(0.45))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY

    panel = slide.shapes.add_shape(1, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.8))
    panel.fill.solid(); panel.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    panel.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    panel.line.width = Pt(1)
    tf = panel.text_frame; tf.word_wrap = True; tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.2)
    for i, line in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.text = line
        par.font.size = Pt(18)
        par.font.color.rgb = TEXT
        par.space_after = Pt(10)

add_title_slide("Multi-Plane Runtime Manager", "Engineering demo: architecture, flow, safety, and extensibility")

add_content_slide(
    "Problem statement",
    [
        "A device-side runtime must handle multiple operation classes under one service boundary.",
        "Not all requests are equally safe: triage, management, control, and config-apply differ in risk.",
        "The runtime must execute commands safely, traceably, and within a tight operational envelope.",
        "It should support both static catalog tools and dynamic plugin expansion without destabilizing the process."
    ],
    accent_color=BLUE,
    note="This is a secure execution engine for diagnostic and operational workflows."
)

add_content_slide(
    "High-level architecture",
    [
        "Cloud or management systems issue WRP requests to the device runtime through Parodus.",
        "The main process owns socket lifecycle, registration, receive loop, and lifecycle handling.",
        "Worker threads isolate per-request execution and response generation.",
        "Catalog and plugin managers resolve tool definitions and dynamic libraries.",
        "Safety and resource checks gate execution before the command is launched."
    ],
    accent_color=NAVY,
    note="Core model: receive → resolve → validate → execute → respond."
)

add_diag_slide(
    "Architecture diagram",
    [
        "Cloud / Management -> Parodus -> Multi-Plane Runtime Manager -> Catalog / Plugin Manager -> Safety Gate -> Command Executor",
        "Result is sent back as WRP response over the PUSH socket to Parodus and then to the cloud system.",
        "Per-plane policy selection determines whether a tool uses a diagnostic-safe or admin-unsafe policy path.",
        "The runtime keeps a plugin manifest and logs for operational traceability and debugging."
    ]
)

add_two_col_slide(
    "Request flow example",
    [
        "Example request: triage plane, tool = device_uptime, command = cat /proc/uptime",
        "Request arrives at PULL socket from Parodus.",
        "Main thread decodes outer WRP and inner payload.",
        "Worker thread looks up the tool in the triage catalog."
    ],
    [
        "The catalog resolves the command and the plane policy class.",
        "The runtime verifies the command against the allowlist and denylist rules.",
        "The process checks CPU, memory, disk, and process thresholds.",
        "A bounded timeout is applied and output is captured before the response is returned."
    ],
    accent_color=ACCENT
)

add_content_slide(
    "Execution safety features",
    [
        "Plane-aware policy classification: triage/management use diagnostic-safe policy; control/config-apply use admin-unsafe policy.",
        "Allowlist and denylist enforcement ensures only approved binaries and command patterns execute.",
        "Resource gate blocks execution when the host is near or over configured thresholds.",
        "Timeout and output caps prevent runaway commands and unbounded output.",
        "Plugin manifests persist loaded tool names and paths instead of relying only on in-memory state."
    ],
    accent_color=GREEN,
    note="Safety is enforced before command launch, not only after execution begins."
)

add_content_slide(
    "Dynamic plugin example",
    [
        "Example plugin: sample_multi_plane_runtime_manager_plugin.c",
        "Exports plugin_echo and plugin_version.",
        "The plugin is discovered from a configured plugin directory and registered as a runtime tool.",
        "A caller can request the plugin tool by plane and tool name; the runtime resolves it and executes it under the applicable policy gate.",
        "If a conflict is detected, the configured plugin conflict policy decides whether to reject or log the duplicate."
    ],
    accent_color=BLUE,
    note="This demonstrates runtime extensibility without rebuilding the main binary."
)

add_content_slide(
    "Configuration model",
    [
        "Configuration is defined in rdk-service-hub.conf.",
        "Environment variables override config file values at runtime.",
        "Catalog, plugin, resource threshold, allowlist, timeout, and log paths are all configurable.",
        "This allows environment-specific deployment tuning while preserving consistent runtime behavior.",
        "Example settings include plugin path, timeout defaults, and per-plane policy selection keys."
    ],
    accent_color=NAVY,
    note="Run-time flexibility is intentionally separated from core execution logic."
)

add_content_slide(
    "Demo outcomes",
    [
        "The runtime resolves a request to the correct plane without mixing tool classes across operational boundaries.",
        "Safe read-only tools can run under a diagnostic policy path while higher-risk actions stay behind stricter gates.",
        "Plugins can be added and discovered without a fresh rebuild of the main service.",
        "The system delivers bounded execution, observability, and a deterministic response contract for callers."
    ],
    accent_color=GREEN,
    note="Engineering demo summary: controlled execution with extensibility and clear operational boundaries."
)

prs.save(out_file)
print(f"Created presentation: {out_file}")
