from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

out_file = "/Users/vpalac201@cable.comcast.com/Project/rdk-service-hub/docs/rdk-service-hub-brief.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BLUE = RGBColor(0x1F, 0x6F, 0xB6)
TEXT = RGBColor(0x1E, 0x2A, 0x39)
ACCENT = RGBColor(0x27, 0x9E, 0xC2)
LIGHT = RGBColor(0xF7, 0xF9, 0xFC)
GREEN = RGBColor(0x20, 0x9B, 0x5A)


def add_title_slide(title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT

    strip = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.22))
    strip.fill.solid(); strip.fill.fore_color.rgb = NAVY
    strip.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(8.3), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.75), Inches(2.0), Inches(11.2), Inches(0.8))
        stf = sub.text_frame
        sp = stf.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(16)
        sp.font.color.rgb = TEXT

    card = slide.shapes.add_shape(1, Inches(9.3), Inches(1.6), Inches(2.9), Inches(1.8))
    card.fill.solid(); card.fill.fore_color.rgb = BLUE
    card.line.fill.background()
    card.text = "Runtime\nManager"
    card.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    card.text_frame.paragraphs[0].font.size = Pt(20)
    card.text_frame.paragraphs[0].font.bold = True
    card.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def add_content_slide(title, bullets, accent_color=BLUE, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT

    strip = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.55))
    strip.fill.solid(); strip.fill.fore_color.rgb = accent_color
    strip.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.8), Inches(7.5), Inches(0.45))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY

    panel = slide.shapes.add_shape(1, Inches(0.7), Inches(1.5), Inches(11.8), Inches(4.7))
    panel.fill.solid(); panel.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    panel.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    panel.line.width = Pt(1)

    tf = panel.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.15); tf.margin_bottom = Inches(0.15)

    for i, bullet in enumerate(bullets):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.text = bullet
        par.level = 0
        par.bullet = True
        par.font.size = Pt(20)
        par.font.color.rgb = TEXT
        par.space_after = Pt(8)

    if note:
        note_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.45))
        np = note_box.text_frame.paragraphs[0]
        np.text = note
        np.font.size = Pt(11)
        np.font.color.rgb = RGBColor(0x4A, 0x5A, 0x68)


def add_two_col_slide(title, left, right, accent_color=BLUE):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT

    strip = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.55))
    strip.fill.solid(); strip.fill.fore_color.rgb = accent_color
    strip.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.8), Inches(8.0), Inches(0.45))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY

    left_box = slide.shapes.add_shape(1, Inches(0.7), Inches(1.5), Inches(5.5), Inches(4.9))
    left_box.fill.solid(); left_box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    left_box.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    left_box.line.width = Pt(1)
    ltf = left_box.text_frame
    ltf.word_wrap = True
    ltf.margin_left = Inches(0.25); ltf.margin_right = Inches(0.2)
    for i, item in enumerate(left):
        par = ltf.paragraphs[0] if i == 0 else ltf.add_paragraph()
        par.text = item
        par.bullet = True
        par.font.size = Pt(18)
        par.font.color.rgb = TEXT
        par.space_after = Pt(8)

    right_box = slide.shapes.add_shape(1, Inches(6.7), Inches(1.5), Inches(5.8), Inches(4.9))
    right_box.fill.solid(); right_box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    right_box.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    right_box.line.width = Pt(1)
    rtf = right_box.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.25); rtf.margin_right = Inches(0.2)
    for i, item in enumerate(right):
        par = rtf.paragraphs[0] if i == 0 else rtf.add_paragraph()
        par.text = item
        par.bullet = True
        par.font.size = Pt(18)
        par.font.color.rgb = TEXT
        par.space_after = Pt(8)


def add_flow_slide(title, lines):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT

    strip = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.55))
    strip.fill.solid(); strip.fill.fore_color.rgb = ACCENT
    strip.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.8), Inches(10), Inches(0.45))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY

    panel = slide.shapes.add_shape(1, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.8))
    panel.fill.solid(); panel.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    panel.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    panel.line.width = Pt(1)

    tf = panel.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.2)
    for i, line in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.text = line
        par.font.size = Pt(18)
        par.font.color.rgb = TEXT
        par.space_after = Pt(10)


# Executive deck
add_title_slide(
    "Multi-Plane Runtime Manager",
    "A safe, extensible runtime for plane-aware diagnostics and command execution"
)

add_content_slide(
    "What the module does",
    [
        "Receives device-side diagnostic and management requests from cloud or management systems.",
        "Routes each request by plane: triage, management, control, or config-apply.",
        "Resolves tools from JSON catalogs or optional dynamic .so plugins.",
        "Enforces command safety, resource gating, and timeout protection before execution.",
        "Returns structured results with exit status, output, and operational logs."
    ],
    accent_color=BLUE,
    note="Purpose: provide controlled execution without exposing unsafe command paths."
)

add_two_col_slide(
    "Architecture at a glance",
    [
        "Cloud / management layer submits requests and receives responses.",
        "Parodus provides the transport bridge into the device runtime.",
        "Main process handles socket lifecycle, registration, and request intake.",
        "Worker threads isolate per-request execution and response generation."
    ],
    [
        "Catalog manager resolves tool definitions and plane metadata.",
        "Safety layer checks allowlists, denylist rules, and resource thresholds.",
        "Executor runs commands with constrained argv-based invocation.",
        "Manifest and logs provide visibility into loaded tools and runtime state."
    ],
    accent_color=NAVY
)

add_flow_slide(
    "Execution workflow",
    [
        "1. Incoming WRP request arrives from Parodus over the PULL socket.",
        "2. Runtime identifies plane, tool name, and any override parameters.",
        "3. Tool is resolved from catalog or plugin registry.",
        "4. Safety policy validates command, resource readiness, and execution limits.",
        "5. Command runs with bounded output capture and timeout protection.",
        "6. Response is packaged and sent back to the caller with execution result."
    ]
)

add_content_slide(
    "Key capabilities and support",
    [
        "Plane-aware policy splits for diagnostic-safe and admin-unsafe command classes.",
        "Dynamic plugin loading for .so-based extension without changing the core binary.",
        "Resource guardrails for CPU, memory, disk, and process stress conditions.",
        "Timeout, loop, and output safeguards to prevent runaway or unbounded commands.",
        "Deployment tuning via config file and environment-variable overrides."
    ],
    accent_color=GREEN,
    note="Supported deployment model: safe extenders for diagnostics and operational management."
)

add_content_slide(
    "Module support areas",
    [
        "Diagnostics and health validation for device-side operations.",
        "Operational management for status, maintenance, and health checks.",
        "Controlled command handling for system operations with policy enforcement.",
        "Config-change workflows where stricter guardrails are required.",
        "Extensible runtime for future tool and plugin additions without core rework."
    ],
    accent_color=ACCENT,
    note="This runtime is designed for embedded device operations and managed fleet diagnostics."
)

add_content_slide(
    "Operational summary",
    [
        "The runtime combines routing, catalog lookup, dynamic plugin execution, and safety enforcement in one process.",
        "It is suitable for diagnostic workloads, controlled device management, and safe tool execution under policy constraints.",
        "The platform is designed for observability, enforceable command safety, extensibility, and deployment customization.",
        "Operational teams can tune policy and thresholds without changing the core execution model."
    ],
    accent_color=GREEN,
    note="Outcome: a flexible but disciplined runtime for multi-plane device management and diagnostics."
)

prs.save(out_file)
print(f"Created presentation: {out_file}")
