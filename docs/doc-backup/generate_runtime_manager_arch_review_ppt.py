from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

out_file = "/Users/vpalac201@cable.comcast.com/Project/rdk-service-hub/docs/rdk-service-hub-arch-review.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BLUE = RGBColor(0x1F, 0x6F, 0xB6)
TEXT = RGBColor(0x1E, 0x2A, 0x39)
ACCENT = RGBColor(0x27, 0x9E, 0xC2)
LIGHT = RGBColor(0xF7, 0xF9, 0xFC)
GREEN = RGBColor(0x20, 0x9B, 0x5A)
ORANGE = RGBColor(0xD9, 0x7A, 0x2A)


def add_title_slide(title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
    strip = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.22))
    strip.fill.solid(); strip.fill.fore_color.rgb = NAVY
    strip.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(8.5), Inches(0.9))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY

    sub = slide.shapes.add_textbox(Inches(0.75), Inches(2.1), Inches(11.2), Inches(0.8))
    sp = sub.text_frame.paragraphs[0]
    sp.text = subtitle
    sp.font.size = Pt(16)
    sp.font.color.rgb = TEXT

    card = slide.shapes.add_shape(1, Inches(9.2), Inches(1.6), Inches(3.0), Inches(1.8))
    card.fill.solid(); card.fill.fore_color.rgb = BLUE
    card.line.fill.background(); card.text = "Review\nBoard"
    card.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    card.text_frame.paragraphs[0].font.size = Pt(18)
    card.text_frame.paragraphs[0].font.bold = True
    card.text_frame.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def add_content_slide(title, bullets, accent_color=BLUE, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
    strip = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.55))
    strip.fill.solid(); strip.fill.fore_color.rgb = accent_color
    strip.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.8), Inches(8.0), Inches(0.45))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY

    panel = slide.shapes.add_shape(1, Inches(0.7), Inches(1.5), Inches(11.7), Inches(4.8))
    panel.fill.solid(); panel.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    panel.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    panel.line.width = Pt(1)
    tf = panel.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.15); tf.margin_bottom = Inches(0.15)

    for i, bullet in enumerate(bullets):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.text = bullet
        par.level = 0
        par.bullet = True
        par.font.size = Pt(19)
        par.font.color.rgb = TEXT
        par.space_after = Pt(8)

    if note:
        note_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.45))
        np = note_box.text_frame.paragraphs[0]
        np.text = note
        np.font.size = Pt(11)
        np.font.color.rgb = RGBColor(0x4A, 0x5A, 0x68)


def add_two_col_slide(title, left, right, accent_color=BLUE):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
    strip = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.55))
    strip.fill.solid(); strip.fill.fore_color.rgb = accent_color
    strip.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.8), Inches(8.0), Inches(0.45))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY

    left_box = slide.shapes.add_shape(1, Inches(0.7), Inches(1.5), Inches(5.4), Inches(4.9))
    left_box.fill.solid(); left_box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    left_box.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    left_box.line.width = Pt(1)
    ltf = left_box.text_frame; ltf.word_wrap = True
    ltf.margin_left = Inches(0.25); ltf.margin_right = Inches(0.2)
    for i, item in enumerate(left):
        par = ltf.paragraphs[0] if i == 0 else ltf.add_paragraph()
        par.text = item
        par.bullet = True
        par.font.size = Pt(18)
        par.font.color.rgb = TEXT
        par.space_after = Pt(8)

    right_box = slide.shapes.add_shape(1, Inches(6.7), Inches(1.5), Inches(5.9), Inches(4.9))
    right_box.fill.solid(); right_box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    right_box.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    right_box.line.width = Pt(1)
    rtf = right_box.text_frame; rtf.word_wrap = True
    rtf.margin_left = Inches(0.25); rtf.margin_right = Inches(0.2)
    for i, item in enumerate(right):
        par = rtf.paragraphs[0] if i == 0 else rtf.add_paragraph()
        par.text = item
        par.bullet = True
        par.font.size = Pt(18)
        par.font.color.rgb = TEXT
        par.space_after = Pt(8)


def add_flow_slide(title, lines):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = LIGHT
    strip = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.55))
    strip.fill.solid(); strip.fill.fore_color.rgb = ACCENT
    strip.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.8), Inches(9.0), Inches(0.45))
    p = title_box.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY

    panel = slide.shapes.add_shape(1, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.9))
    panel.fill.solid(); panel.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    panel.line.color.rgb = RGBColor(0xD8, 0xE2, 0xEC)
    panel.line.width = Pt(1)
    tf = panel.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.2)
    for i, line in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.text = line
        par.font.size = Pt(18)
        par.font.color.rgb = TEXT
        par.space_after = Pt(10)


add_title_slide(
    "Multi-Plane Runtime Manager",
    "Internal architecture review: interfaces, control plane, runtime boundaries, and safety model"
)

add_content_slide(
    "Review goals",
    [
        "Confirm a clear runtime boundary between cloud-facing ingress, internal execution, and tool safety policy.",
        "Verify safe routing across triage, management, control, and config-apply workloads.",
        "Assess plugin extensibility and catalog-driven execution against operational risk.",
        "Review runtime observability, configuration precedence, and policy tunability.",
        "Validate the design against failure modes: timeouts, loops, blocked commands, and resource exhaustion."
    ],
    accent_color=BLUE,
    note="Architecture review focus: separation of concerns, fail-safe execution, and auditability."
)

add_two_col_slide(
    "System boundaries",
    [
        "Cloud / management side submits requests and receives responses.",
        "Parodus provides the device-side transport endpoint and control bridge.",
        "Main runtime thread owns lifecycle, receive loop, registration, and socket management.",
        "Worker threads handle execution and response generation independently."
    ],
    [
        "Catalog manager interprets static tool definitions and per-plane metadata.",
        "Plugin manager loads and tracks .so-based tools from configured directories.",
        "Safety gate enforces allowlist, denylist, and resource checks before execution.",
        "Executor reports output, exit code, logs, and policy outcomes back to the caller."
    ],
    accent_color=NAVY
)

add_flow_slide(
    "Runtime control flow",
    [
        "1. Request enters through the PULL socket from Parodus.",
        "2. Main thread decodes WRP outer and inner payload metadata.",
        "3. Request is dispatched to a worker thread for per-request execution context.",
        "4. Catalog or plugin registry resolves the tool name and plane mapping.",
        "5. Policy validation checks command safety, resource readiness, and timeout bounds.",
        "6. Command executes under a bounded output and time budget, then a structured response is returned."
    ]
)

add_content_slide(
    "Safety model",
    [
        "Plane-aware policy selection separates diagnostic-safe and admin-unsafe classes.",
        "Allowlist and denylist enforcement is applied before any command is launched.",
        "The runtime validates resource limits for CPU, memory, disk, and process pressure.",
        "Timeout, loop, and output-cap controls reduce risk of process runaway or output flooding.",
        "Failed or blocked requests return clear operational messages instead of silently falling through."
    ],
    accent_color=GREEN,
    note="The model aims to fail closed where command safety cannot be guaranteed."
)

add_two_col_slide(
    "Extensibility review",
    [
        "Static catalogs provide a stable, packaged tool surface for deployment.",
        "Dynamic plugins allow optional extension without rebuilding the main runtime.",
        "Per-plane plugin directories isolate diagnostic and higher-risk tool sets.",
        "Manifest persistence tracks loaded tool names and .so paths for observability."
    ],
    [
        "Conflict policy prevents duplicate tool registration from silently overriding official entries.",
        "Environment variables override config defaults at runtime without code changes.",
        "Logging and startup feature tables improve operational validation during deployment.",
        "This keeps the system adaptive while preserving a well-defined runtime contract."
    ],
    accent_color=ACCENT
)

add_content_slide(
    "Architecture concerns and mitigations",
    [
        "Concern: execution policy drift between planes -> Mitigation: explicit per-plane policy class selection.",
        "Concern: destructive command invocation -> Mitigation: allowlist + denylist + no-shell execution path.",
        "Concern: host overload -> Mitigation: resource gate checks before command launch.",
        "Concern: runaway processes -> Mitigation: timeout, output cap, and loop detection guardrails.",
        "Concern: dynamic tool sprawl -> Mitigation: plugin conflict policy, manifest tracking, and log visibility."
    ],
    accent_color=ORANGE,
    note="The overall design remains intentionally conservative around command execution."
)

add_content_slide(
    "Board recommendation",
    [
        "Proceed with the runtime model as a controlled, plane-aware execution framework.",
        "Keep policy separation explicit and enforce command safety at the execution boundary.",
        "Continue to validate plugin and catalog behavior under real deployment conditions.",
        "Maintain a persistent manifest and clear log coverage as the operational audit trail.",
        "Treat dynamic extension as valuable, but gate it behind policy and conflict management."
    ],
    accent_color=GREEN,
    note="Recommendation: approve the architectural pattern while continuing focused validation and hardening."
)

prs.save(out_file)
print(f"Created presentation: {out_file}")
